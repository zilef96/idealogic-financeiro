import openpyxl, sys
XLSX="docs/resources/Orçamento 2026.xlsx"
OUT="prisma/seed-orcado-mensal.sql"
itens=set(int(x) for x in open("/tmp/item_codigos.txt").read().split())
wb=openpyxl.load_workbook(XLSX, data_only=True)
ws=wb["Execução Orçamentária"]
MES={"janeiro":1,"fevereiro":2,"março":3,"marco":3,"abril":4,"maio":5,"junho":6,
     "julho":7,"agosto":8,"setembro":9,"outubro":10,"novembro":11,"dezembro":12}
mescol={}
for c in range(1,ws.max_column+1):
    v=ws.cell(18,c).value
    if isinstance(v,str) and v.strip().lower() in MES:
        mescol[MES[v.strip().lower()]]=c   # coluna = orçado daquele mês
assert len(mescol)==12, mescol
linhas=[]; itens_vistos=set()
for r in range(19, ws.max_row+1):
    b=ws.cell(r,2).value
    if b is None: continue
    try: cod=int(b)
    except (ValueError,TypeError): continue
    if cod not in itens: continue
    itens_vistos.add(cod)
    for mes,col in sorted(mescol.items()):
        v=ws.cell(r,col).value
        v=0.0 if v is None else float(v)
        linhas.append((cod,mes,round(v,2)))
with open(OUT,"w") as f:
    f.write("-- Orçado mensal por item, importado de 'Execução Orçamentária' (Orçamento 2026.xlsx)\n")
    f.write("-- Popula lancamento_realizado.valor_orcado (orçado mensal ajustável) por item x competência.\n")
    f.write("WITH dados(codigo, mes, valor) AS (VALUES\n")
    f.write(",\n".join(f"  ({c}, {m}, {v})" for c,m,v in linhas))
    f.write("\n)\nINSERT INTO lancamento_realizado (conta_item_id, exercicio_id, competencia, valor_orcado)\n")
    f.write("SELECT ci.id, e.id, make_date(2026, d.mes, 1), d.valor::numeric\n")
    f.write("FROM dados d JOIN exercicio e ON e.ano=2026\n")
    f.write("JOIN conta_grupo cg ON cg.exercicio_id=e.id\n")
    f.write("JOIN conta_item ci ON ci.grupo_id=cg.id AND ci.codigo=d.codigo\n")
    f.write("ON CONFLICT (conta_item_id, competencia) DO UPDATE SET valor_orcado=EXCLUDED.valor_orcado, updated_at=now();\n")
print(f"itens casados: {len(itens_vistos)}/{len(itens)} | linhas (item x mes): {len(linhas)}")
print("colunas de mês:", mescol)
